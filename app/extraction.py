"""Document text extraction — the deterministic half of resume ingestion.

Step [0] of the pipeline, before the Profiler agent ever runs:

    bytes  ──▶  THIS MODULE  ──▶  plain text  ──▶  [1] Profiler  ──▶  ProfileDraft
                (libraries)                          (Haiku)

This module answers "what words are on the page, in what order". It has no idea
what an employer is. The model does the semantic half — mapping words onto the
career schema — because "is this line a job title or a project name" is a
judgement call, and D10 puts judgement in the agent layer.

Ported from Caliber `api/src/caliber/extraction.py`. Keeping extraction
deterministic buys four things the model cannot:

  - it runs with no credential and no network, so every test stays offline;
  - it costs nothing, where PDF-as-vision runs ~1.5-3k tokens per page and
    resume parsing is the highest-volume model call in the product;
  - it is provider-independent (D10) — text works behind any model;
  - when a parse comes out wrong, the extracted text shows immediately whether
    extraction garbled it or the model misjudged it. One black box, not two.

WHY NOT pypdf: pypdf reads a page in raw content-stream order, which on a
two-column resume interleaves the skills sidebar into the job history — "Kafka"
lands in the middle of a job title. Measured on this machine: both pypdf and
plain pdfplumber interleaved a two-column resume; the gutter detection below
recovered 100% of the text in reading order, in ~43 ms.

Deviations from Caliber's module are marked DEVIATION and listed in RULEBOOK §16.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass

import pdfplumber

from app.policy import MAX_RESUME_CHARS

log = logging.getLogger(__name__)

# Below this, a document is treated as having no usable text layer — almost
# always a scan or an image-only export. The caller turns that into an
# actionable message rather than handing the model an empty string.
MIN_USEFUL_CHARS = 100

# Refuse before parsing: a file this large is never a resume.
MAX_UPLOAD_BYTES = 8 * 1024 * 1024

# Column detection tuning. Deliberately conservative: a WRONG column split is
# far more damaging than a missed one, because it reorders text that was
# already correct. When in doubt, fall through to plain top-to-bottom.
# These are page geometry, not profiling policy, so they live here and not in
# app/policy.py — nothing here is stamped into a CandidateProfile.
_GUTTER_SCAN = range(28, 73, 2)       # candidate gutters, % of page width
_GUTTER_PAD = 3.0                     # pt of clearance required either side
_MIN_SIDE_SHARE = 0.15                # each column needs >=15% of the words
_MIN_SIDE_WORDS = 8
_MAX_HEADER_FRACTION = 0.40           # a full-width header may occupy the top 40%
_LINE_EPSILON = 1.0                   # pt of slack so a closed line is not re-cropped

# DEVIATION 1 (RULEBOOK §16 O18) — not in Caliber. Caliber accepts ANY clear
# vertical channel as a gutter, with no minimum width, so on a single-column
# page an accidental alignment of inter-word spaces wins and words straddling
# the fake gutter are emitted TWICE, in halves ("...flaky-test quarantin" /
# "e bot; ..."). Measured clearances: a true two-column gap is 41.4pt, the two
# false positives were 3.5pt and 3.9pt. Anything narrower is an accident.
_MIN_CLEARANCE = 12.0

# DEVIATION 4 (O21) — not in Caliber. Most PDFs store no space characters at all;
# pdfplumber infers a word break when the horizontal gap between two glyphs
# exceeds `x_tolerance` (default 3 pt). With a tight font that threshold is wider
# than a real space, so words silently fuse — "BachelorofEngineering",
# "UniversityofPune" — and nothing raises: the text still passes `is_usable` and
# every other check here. It only shows up downstream, where the Profiler
# verifies each model claim as a VERBATIM substring of this text, so a
# space-mangled source rejects every claim.
#
# Measured on one real (git-ignored) resume, 4 probe phrases, `extract_text`:
#
#     setting                 chars   [a-z][A-Z] joins   probes found
#     default (x_tolerance=3)  4126        55              0/4
#     x_tolerance_ratio=0.15   4427        16              4/4
#     x_tolerance=1.5          4427        16              4/4
#
# The remaining 16 joins are real CamelCase (JavaScript, MySQL, GraphQL).
# Ratio over a flat tolerance: the ratio scales the threshold with each glyph's
# font size (`ratio * size`), so an 18 pt heading and 9 pt body get the same
# treatment, where a flat 1.5 pt would over-split the heading. 0.15 sits in the
# middle of the 0.12-0.20 band that gave identical output on that file, so it is
# not tuned to one document. Applied to EVERY pdfplumber call below — word
# boxes, plain text and cropped columns — or the column path would disagree with
# the plain path on where words begin.
_TOLERANCE: dict[str, float] = {"x_tolerance_ratio": 0.15}

# DEVIATION 2 (O19) — not in Caliber. pdfminer emits "(cid:NNN)" for any glyph
# it cannot map to a character; Symbol/Wingdings bullets hit this constantly.
# Left alone it reaches the model as junk tokens inside every bullet line.
_CID = re.compile(r"\(cid:\d+\)")

# Control characters that survive some PDF text layers. NUL in particular breaks
# downstream string handling, and none of them mean anything in a resume.
_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class ExtractionError(ValueError):
    """Base for every unreadable-document fact.

    A ValueError, so a caller written against Caliber's contract
    (``except ValueError``) keeps working.
    """


class UnsupportedFileType(ExtractionError):
    """Not a PDF or DOCX. HTTP 415."""


class FileTooLarge(ExtractionError):
    """Larger than MAX_UPLOAD_BYTES. HTTP 413."""


class UnreadableDocument(ExtractionError):
    """Corrupt, truncated, or not really the type its name claims. HTTP 400."""


class NoTextLayer(ExtractionError):
    """Parsed fine but carries no usable text — a scan or a photo. HTTP 422.

    This is the one case that must never reach the model: an empty string costs
    tokens and comes back as a confidently invented profile.
    """


@dataclass(frozen=True)
class Extraction:
    text: str
    pages: int
    layout: str  # "single-column" | "multi-column" | "docx" | "empty"

    @property
    def is_usable(self) -> bool:
        return len(self.text.strip()) >= MIN_USEFUL_CHARS


@dataclass(frozen=True)
class ResumeSource:
    """What the Profiler is actually given, plus what it cost to get there."""

    text: str                 # clipped to MAX_RESUME_CHARS, ready for run_profiler()
    extraction: Extraction    # the full extraction, unclipped
    clipped_chars: int        # 0 unless the document was longer than the cap

    @property
    def was_clipped(self) -> bool:
        return self.clipped_chars > 0


def _clean(text: str) -> str:
    """Normalise whitespace without destroying line structure.

    Line breaks carry meaning in a resume (they separate roles from bullets), so
    they survive; runs of blank lines and trailing spaces do not.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CID.sub("•", text)      # DEVIATION 2
    text = _CONTROLS.sub("", text)       # DEVIATION 2
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _find_gutter(words: list[dict], width: float, height: float) -> tuple[float, float] | None:
    """Locate a vertical gutter separating two text columns.

    Returns ``(gutter_x, header_cut_y)`` or None.

    ``header_cut_y`` is where the columns begin: a resume usually has a
    full-width name/contact header above a two-column body, and splitting the
    page from y=0 would tear that header in half. Text above the cut is read
    full width.

    A candidate x is a gutter when no word crosses it below the cut, both sides
    carry a real share of the words, and the clear channel is wide enough to be
    a column gap rather than a run of aligned spaces.
    """
    if not words:
        return None

    best: tuple[float, float, float] | None = None  # (clearance, x, cut)

    for pct in _GUTTER_SCAN:
        x = width * pct / 100.0
        straddlers = [w for w in words if w["x0"] < x - _GUTTER_PAD and w["x1"] > x + _GUTTER_PAD]

        if straddlers:
            # Columns may still start below whatever crosses this line.
            cut = max(w["bottom"] for w in straddlers)
            # Close the line. The cut lands on the lowest STRADDLING word, but
            # its neighbours on the same visual line sit a fraction lower (font
            # metrics differ per glyph run). Cropping at the raw cut then slices
            # that line down the middle and emits it three times: once in the
            # header, then a fragment in each column.
            for _ in range(5):
                spanning = [w for w in words if w["top"] < cut - 0.5 and w["bottom"] > cut]
                if not spanning:
                    break
                cut = max(w["bottom"] for w in spanning)
            cut += _LINE_EPSILON
            if cut > height * _MAX_HEADER_FRACTION:
                continue  # something crosses too far down; not a two-column body
        else:
            cut = 0.0

        body = [w for w in words if w["top"] >= cut]
        if not body:
            continue
        left = [w for w in body if w["x1"] <= x]
        right = [w for w in body if w["x0"] >= x]
        threshold = max(_MIN_SIDE_WORDS, int(len(body) * _MIN_SIDE_SHARE))
        if len(left) < threshold or len(right) < threshold:
            continue

        # Prefer the gutter with the widest clear channel — the true column gap
        # rather than an incidental alignment inside one column.
        clearance = min(
            x - max((w["x1"] for w in left), default=x),
            min((w["x0"] for w in right), default=x) - x,
        )
        if clearance < _MIN_CLEARANCE:  # DEVIATION 1
            continue
        if best is None or clearance > best[0]:
            best = (clearance, x, cut)

    return (best[1], best[2]) if best else None


def _page_text(page) -> tuple[str, bool]:
    """Text for one page, column-aware. Returns ``(text, was_multi_column)``."""
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False, **_TOLERANCE)
    gutter = _find_gutter(words, page.width, page.height)

    if gutter is None:
        return (page.extract_text(**_TOLERANCE) or ""), False

    x, cut = gutter
    parts: list[str] = []

    # Full-width header above the columns, if any.
    if cut > 0:
        header = page.crop((0, 0, page.width, cut)).extract_text(**_TOLERANCE)
        if header:
            parts.append(header)

    # Left column, then right — the order a human reads them.
    for box in ((0, cut, x, page.height), (x, cut, page.width, page.height)):
        col = page.crop(box).extract_text(**_TOLERANCE)
        if col:
            parts.append(col)

    return "\n".join(parts), True


def _reading_failure(kind: str) -> UnreadableDocument:
    # pdfplumber and python-docx raise library-specific errors on corrupt or
    # mislabeled files (PdfminerException, BadZipFile, PackageNotFoundError…) —
    # none of them ValueError. At this boundary every unreadable file is the
    # same domain fact, and the message is already safe to show a candidate.
    return UnreadableDocument(
        f"Could not read that {kind} - the file looks corrupt, or is not really a {kind}."
    )


def extract_pdf(raw: bytes) -> Extraction:
    multi = False
    chunks: list[str] = []
    try:
        pdf_cm = pdfplumber.open(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001 - library types, normalised at this boundary
        raise _reading_failure("PDF") from exc
    with pdf_cm as pdf:
        for page in pdf.pages:
            try:
                text, was_multi = _page_text(page)
            except Exception as exc:  # noqa: BLE001 - one bad page must not lose the rest
                log.warning("pdf page %s failed: %s", page.page_number, exc)
                text, was_multi = (page.extract_text(**_TOLERANCE) or ""), False
            multi = multi or was_multi
            if text:
                chunks.append(text)
        pages = len(pdf.pages)

    text = _clean("\n\n".join(chunks))
    layout = "empty" if not text else ("multi-column" if multi else "single-column")
    return Extraction(text=text, pages=pages, layout=layout)


def _row_text(row) -> str:
    """One table row as a line.

    A table used for layout repeats the same cell across a merged row, so
    de-duplicate before joining.
    """
    seen: list[str] = []
    for cell in row.cells:
        value = cell.text.strip().replace("\n", " ")
        if value and value not in seen:
            seen.append(value)
    return " | ".join(seen) if len(seen) > 1 else (seen[0] if seen else "")


def _docx_blocks(doc):
    """Yield paragraphs and tables in document-body order.

    DEVIATION 3 (O20) — Caliber walks ``doc.paragraphs`` fully and only then
    ``doc.tables``, which lands an Experience table at the very bottom of the
    text, far from its own "Experience" heading. Readable, but the model then
    has to re-associate them. Walking the body element keeps each table under
    the heading that introduces it. Falls back to Caliber's order if the
    python-docx internals ever move.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        body = doc.element.body
    except AttributeError:  # pragma: no cover - defensive
        yield from doc.paragraphs
        yield from doc.tables
        return

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def extract_docx(raw: bytes) -> Extraction:
    """DOCX needs no layout analysis — the XML already carries reading order.

    It DOES need tables. Many resume templates lay the whole document out in a
    table, and reading only ``document.paragraphs`` returns nothing at all for
    those files: the text lives in cell paragraphs, which are not in that list.
    """
    import docx
    from docx.table import Table

    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001 - BadZipFile/PackageNotFoundError are library types
        raise _reading_failure("DOCX") from exc

    parts: list[str] = []
    for block in _docx_blocks(doc):
        if isinstance(block, Table):
            for row in block.rows:
                line = _row_text(row)
                if line:
                    parts.append(line)
        elif block.text.strip():
            parts.append(block.text)

    text = _clean("\n".join(parts))
    return Extraction(text=text, pages=1, layout="empty" if not text else "docx")


def extract(filename: str, raw: bytes) -> Extraction:
    """Dispatch on extension.

    Every failure is an :class:`ExtractionError`. This module knows nothing
    about HTTP, but each subclass maps to exactly one status — see the class
    docstrings.
    """
    if len(raw) > MAX_UPLOAD_BYTES:
        raise FileTooLarge(
            f"That file is {len(raw) / 1_048_576:.1f} MB; the limit is "
            f"{MAX_UPLOAD_BYTES // 1_048_576} MB."
        )
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return extract_pdf(raw)
    if name.endswith(".docx"):
        return extract_docx(raw)
    raise UnsupportedFileType("Unsupported file type - upload a PDF or DOCX.")


def load_resume_text(filename: str, raw: bytes) -> ResumeSource:
    """Step [0] in full: bytes in, Profiler-ready text out.

    Raises an :class:`ExtractionError` for every case the model must never see.
    Clipping is reported rather than silent, so the caller can tell the
    candidate that only the first ``MAX_RESUME_CHARS`` characters were read.
    """
    extraction = extract(filename, raw)
    if not extraction.is_usable:
        raise NoTextLayer(
            "That file has no text layer - it looks like a scan or a photo. "
            "Upload the original file, export it from Word or Google Docs as PDF, "
            "or paste your resume as text."
        )
    text = extraction.text
    clipped = max(0, len(text) - MAX_RESUME_CHARS)
    return ResumeSource(text=text[:MAX_RESUME_CHARS], extraction=extraction, clipped_chars=clipped)
